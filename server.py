from flask import Flask, request, jsonify, send_from_directory
from openpyxl import load_workbook
from flask_cors import CORS
import os

app = Flask(__name__, static_folder='.')
CORS(app)

EXCEL_FILE = 'alumnos.xlsx'

# Crear archivo Excel si no existe
def crear_archivo_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        # Crear encabezado
        ws.append(["Nombre", "Evaluador"])
        wb.save(EXCEL_FILE)

# Llamada a la función para crear el archivo si no existe
crear_archivo_excel()

@app.route('/')
def index():
    return send_from_directory('.', 'index.html')

@app.route('/script.js')
def js():
    return send_from_directory('.', 'script.js')

@app.route('/style.css')
def css():
    return send_from_directory('.', 'style.css')

@app.route('/nombres')
def obtener_nombres():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    nombres = [cell.value for cell in ws['A'][1:] if cell.value]  # Salta encabezado
    return jsonify(nombres)

@app.route('/guardar', methods=['POST'])
def guardar_respuestas():
    data = request.json
    evaluador = data['evaluador']
    respuestas = data['respuestas']

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Crear encabezados si no existen
    headers = [cell.value for cell in ws[1]]
    if "Evaluaciones Recibidas" not in headers:
        for nombre in respuestas:
            if nombre not in headers:
                ws.cell(row=1, column=len(headers) + 1).value = nombre
                headers.append(nombre)

    # Guardar las evaluaciones
    for nombre_evaluado, puntaje in respuestas.items():
        fila = None
        columna = None

        # Buscar fila (nombre evaluado)
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == nombre_evaluado:
                fila = row
                break

        # Buscar columna (nombre del evaluador o pregunta)
        for col in range(2, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == evaluador:
                columna = col
                break

        # Si columna para este evaluador no existe, la crea
        if columna is None:
            columna = ws.max_column + 1
            ws.cell(row=1, column=columna).value = evaluador

        if fila is not None:
            ws.cell(row=fila, column=columna).value = puntaje

    wb.save(EXCEL_FILE)
    return jsonify({"mensaje": "Respuestas guardadas correctamente."})

if __name__ == '__main__':
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)
